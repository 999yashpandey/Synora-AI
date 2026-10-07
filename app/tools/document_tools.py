from pathlib import Path

from docx import Document

from app.security.permissions import PermissionManager
from app.tools.file_tools import SPECIAL_FOLDERS


permissions = PermissionManager()


def create_word_document(
    filename: str,
    content: str,
    location: str = "desktop",
):
    """
    Create a .docx document in an approved Synora folder.

    Args:
        filename: Desired document filename.
        content: Text to write into the document.
        location: Approved special folder such as desktop/documents.

    Returns:
        Human-readable result string.
    """

    location = location.lower().strip()

    if location not in SPECIAL_FOLDERS:
        return (
            f"I don't have access to the '{location}' folder."
        )

    folder = SPECIAL_FOLDERS[location]

    if not folder.exists() or not folder.is_dir():
        return (
            f"I couldn't find the '{location}' folder."
        )

    if not permissions.is_allowed(folder):
        return (
            "I don't have permission to create files "
            "in that folder."
        )

    filename = filename.strip().strip("\"'")

    if not filename:
        return "Please provide a filename."

    if not filename.lower().endswith(".docx"):
        filename += ".docx"

    output_path = folder / filename

    try:
        document = Document()

        document.add_paragraph(content)

        document.save(output_path)

    except Exception as error:
        return (
            f"I couldn't create the Word document: {error}"
        )

    if not output_path.exists():
        return (
            "The document operation finished, "
            "but I couldn't verify that the file was created."
        )

    return (
        f"Created {output_path.name} in your "
        f"{location} folder."
    )


def create_word_document_with_title(
    filename: str,
    title: str,
    content: str,
    location: str = "desktop",
):
    """
    Create a Word document with a title and body content.
    """

    location = location.lower().strip()

    if location not in SPECIAL_FOLDERS:
        return (
            f"I don't have access to the '{location}' folder."
        )

    folder = SPECIAL_FOLDERS[location]

    if not folder.exists() or not folder.is_dir():
        return (
            f"I couldn't find the '{location}' folder."
        )

    if not permissions.is_allowed(folder):
        return (
            "I don't have permission to create files "
            "in that folder."
        )

    filename = filename.strip().strip("\"'")

    if not filename:
        return "Please provide a filename."

    if not filename.lower().endswith(".docx"):
        filename += ".docx"

    output_path = folder / filename

    try:
        document = Document()

        document.add_heading(title, level=1)

        document.add_paragraph(content)

        document.save(output_path)

    except Exception as error:
        return (
            f"I couldn't create the Word document: {error}"
        )

    if not output_path.exists():
        return (
            "The document operation finished, "
            "but I couldn't verify that the file was created."
        )

    return (
        f"Created {output_path.name} in your "
        f"{location} folder."
    )


if __name__ == "__main__":

    print(
        create_word_document_with_title(
            filename="synora_test.docx",
            title="Synora Document Test",
            content=(
                "This document was created by Synora "
                "using its local document tool."
            ),
            location="desktop",
        )
    ) 