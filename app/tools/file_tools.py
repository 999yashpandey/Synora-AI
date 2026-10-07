import os
from pathlib import Path

from app.security.permissions import PermissionManager


permissions = PermissionManager()

HOME = Path.home()


# ============================================================
# SPECIAL FOLDERS
# ============================================================

def get_special_folders():
    folders = {}

    candidates = {
        "desktop": [
            HOME / "Desktop",
            HOME / "OneDrive" / "Desktop",
        ],
        "downloads": [
            HOME / "Downloads",
            HOME / "OneDrive" / "Downloads",
        ],
        "pictures": [
            HOME / "Pictures",
            HOME / "OneDrive" / "Pictures",
        ],
        "documents": [
            HOME / "Documents",
            HOME / "OneDrive" / "Documents",
        ],
    }

    for name, paths in candidates.items():

        for path in paths:

            if path.exists() and path.is_dir():
                folders[name] = path.resolve()
                break

    return folders


SPECIAL_FOLDERS = get_special_folders()


# ============================================================
# PERMISSION
# ============================================================

def check_access(path):
    return permissions.is_allowed(path)


# ============================================================
# OPEN FOLDER
# ============================================================

def open_folder(folder_name: str):

    folder_name = folder_name.lower().strip()

    if folder_name not in SPECIAL_FOLDERS:
        return f"I don't have access to the folder '{folder_name}'."

    path = SPECIAL_FOLDERS[folder_name]

    if not path.exists():
        return f"I couldn't find the folder '{folder_name}'."

    if not check_access(path):
        return "I don't have permission to access that folder."

    try:
        os.startfile(str(path))
        return f"Opening {path.name}."

    except Exception as error:
        return f"I couldn't open {path.name}: {error}"


# ============================================================
# OPEN FILE
# ============================================================

def open_file(file_path: str):

    path = Path(file_path).expanduser()

    if not path.exists():
        return f"I couldn't find '{file_path}'."

    if not path.is_file():
        return f"'{file_path}' is not a file."

    if not check_access(path):
        return (
            "I found the file, but Synora doesn't have "
            "permission to access its location."
        )

    try:
        os.startfile(str(path))
        return f"Opening {path.name}."

    except Exception as error:
        return f"I found {path.name}, but couldn't open it: {error}"


# ============================================================
# FIND FILE
# ============================================================

def find_file(filename: str):

    filename = (
        filename
        .strip()
        .strip("\"'")
        .lower()
    )

    if not filename:
        return None

    filename_stem = Path(filename).stem.lower()

    search_locations = [
        path
        for path in permissions.allowed_directories
        if path.exists() and path.is_dir()
    ]

    # EXACT FILENAME
    for location in search_locations:

        try:

            for item in location.rglob("*"):

                if not item.is_file():
                    continue

                if item.name.lower() == filename:
                    return item

        except (PermissionError, OSError):
            continue

    # FILENAME WITHOUT EXTENSION
    for location in search_locations:

        try:

            for item in location.rglob("*"):

                if not item.is_file():
                    continue

                if item.stem.lower() == filename_stem:
                    return item

        except (PermissionError, OSError):
            continue

    # PARTIAL MATCH
    for location in search_locations:

        try:

            for item in location.rglob("*"):

                if not item.is_file():
                    continue

                if filename_stem in item.stem.lower():
                    return item

        except (PermissionError, OSError):
            continue

    return None


# ============================================================
# OPEN NAMED FILE
# ============================================================

def open_named_file(filename: str):

    file_path = find_file(filename)

    if file_path is None:
        return None

    return open_file(str(file_path))


# ============================================================
# LIST FOLDER
# ============================================================

def list_folder(folder_path: str = "."):

    path = Path(folder_path).expanduser()

    if not path.exists():
        return f"I couldn't find '{folder_path}'."

    if not path.is_dir():
        return f"'{folder_path}' is not a folder."

    if not check_access(path):
        return "I don't have permission to access that folder."

    try:
        items = list(path.iterdir())

    except PermissionError:
        return f"I don't have permission to access {path.name}."

    if not items:
        return f"{path.name} is empty."

    result = []

    for item in items[:30]:

        if item.is_dir():
            result.append(f"[Folder] {item.name}")
        else:
            result.append(f"[File] {item.name}")

    if len(items) > 30:
        result.append(f"...and {len(items) - 30} more items.")

    return "\n".join(result)


# ============================================================
# CREATE TEXT / CODE FILE
# ============================================================

def create_text_file(
    filename: str,
    content: str,
    folder_name: str = "desktop",
):
    """
    Create a text/code file in an approved special folder.

    Used for:
    .py
    .java
    .js
    .ts
    .tsx
    .jsx
    .html
    .css
    .cpp
    .c
    .json
    .md
    .txt
    etc.
    """

    filename = filename.strip().strip("\"'")

    if not filename:
        return "I need a filename."

    folder_name = folder_name.lower().strip()

    if folder_name not in SPECIAL_FOLDERS:
        return (
            f"I don't have access to the folder "
            f"'{folder_name}'."
        )

    directory = SPECIAL_FOLDERS[folder_name]

    if not directory.exists():
        return f"I couldn't find the {folder_name} folder."

    if not check_access(directory):
        return (
            f"I don't have permission to write to "
            f"{directory}."
        )

    path = directory / filename

    try:

        path.write_text(
            content,
            encoding="utf-8",
        )

    except Exception as error:

        return (
            f"I couldn't create {filename}: {error}"
        )

    # IMPORTANT:
    # Verify that the file was actually created.
    if not path.exists() or not path.is_file():

        return (
            f"I couldn't verify creation of "
            f"{filename}."
        )

    return (
        f"Created {path.name} on your "
        f"{directory.name}.\n"
        f"Path: {path}"
    )


# ============================================================
# CREATE WORD DOCUMENT
# ============================================================

def create_docx_file(
    filename: str,
    title: str,
    content: str,
    folder_name: str = "desktop",
):
    """
    Create a real .docx Word document.
    """

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:

        return (
            "The Word document dependency is missing. "
            "Run: pip install python-docx"
        )

    filename = filename.strip().strip("\"'")

    if not filename.lower().endswith(".docx"):
        filename += ".docx"

    folder_name = folder_name.lower().strip()

    if folder_name not in SPECIAL_FOLDERS:
        return (
            f"I don't have access to the folder "
            f"'{folder_name}'."
        )

    directory = SPECIAL_FOLDERS[folder_name]

    if not check_access(directory):
        return (
            f"I don't have permission to write to "
            f"{directory}."
        )

    path = directory / filename

    try:

        document = Document()

        # Title
        title_paragraph = document.add_paragraph()

        title_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        title_run = title_paragraph.add_run(title)

        title_run.bold = True
        title_run.font.size = None

        # Body
        paragraphs = [
            paragraph.strip()
            for paragraph in content.split("\n")
            if paragraph.strip()
        ]

        for paragraph in paragraphs:

            document.add_paragraph(paragraph)

        document.save(str(path))

    except Exception as error:

        return (
            f"I couldn't create {filename}: {error}"
        )

    # IMPORTANT:
    # Verify actual filesystem creation.
    if not path.exists() or not path.is_file():

        return (
            f"I couldn't verify creation of "
            f"{filename}."
        )

    return (
        f"Created {path.name} on your "
        f"{directory.name}.\n"
        f"Path: {path}"
    )